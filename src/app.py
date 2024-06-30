from flask import Flask, request, render_template, send_file
from docx import Document
import os
import tempfile
from io import BytesIO
import zipfile

app = Flask(__name__)

def replace_placeholders(doc, data):
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            if key in paragraph.text:
                inline = paragraph.runs
                for i in range(len(inline)):
                    if key in inline[i].text:
                        inline[i].text = inline[i].text.replace(key, value if value else '')
    return doc

def process_documents(template_dir, data, output_dir):
    for template_file in os.listdir(template_dir):
        if template_file.endswith('.docx'):
            template_path = os.path.join(template_dir, template_file)
            output_path = os.path.join(output_dir, template_file)
            
            document = Document(template_path)
            document = replace_placeholders(document, data)
            document.save(output_path)

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        data = {
            '<Applicant>': request.form.get('applicant'),
            '<Age1>': request.form.get('age1'),
            '<Applicants Father>': request.form.get('applicants_father'),
            '<Address1>': request.form.get('address1'),
            '<Co-Applicant>': request.form.get('co_applicant'),
            '<Age2>': request.form.get('age2'),
            '<Co-Applicants Father>': request.form.get('co_applicants_father'),
            '<Co-applicant Address2>': request.form.get('co_applicant_address2'),
            '<Co-Applicant2>': request.form.get('co_applicant2'),
            '<Age3>': request.form.get('age3'),
            '<Co-Applicants Father2>': request.form.get('co_applicants_father2'),
            '<Co-Applicants Address3>': request.form.get('co_applicants_address3'),
            '<Co-Applicant3>': request.form.get('co_applicant3'),
            '<Age4>': request.form.get('age4'),
            '<Co-Applicants Father3>': request.form.get('co_applicants_father3'),
            '<Co-Applicants Address4>': request.form.get('co_applicants_address4'),
            '<Property Address>': request.form.get('property_address'),
            '<Area>': request.form.get('area'),
            '<Survey No>': request.form.get('survey_no'),
            '<List1>': request.form.get('list1'),
            '<List2>': request.form.get('list2'),
            '<List3>': request.form.get('list3'),
            '<List4>': request.form.get('list4'),
            '<List5>': request.form.get('list5'),
            '<List6>': request.form.get('list6'),
            '<North>': request.form.get('north'),
            '<South>': request.form.get('south'),
            '<East>': request.form.get('east'),
            '<West>': request.form.get('west'),
            '<Loan Amount>': request.form.get('loan_amount'),
            '<Sanction Date>': request.form.get('sanction_date'),
            '<Take Over from>': request.form.get('take_over_from'),
            '<In words>': request.form.get('in_words')
        }

        temp_dir = tempfile.mkdtemp()
        template_dir = 'templates'
        output_dir = os.path.join(temp_dir, 'outputs')
        os.makedirs(output_dir, exist_ok=True)
        
        process_documents(template_dir, data, output_dir)
        
        output_zip = os.path.join(temp_dir, 'documents.zip')
        with zipfile.ZipFile(output_zip, 'w') as zipf:
            for root, _, files in os.walk(output_dir):
                for file in files:
                    zipf.write(os.path.join(root, file), os.path.relpath(os.path.join(root, file), output_dir))
        
        return send_file(output_zip, as_attachment=True, attachment_filename='documents.zip')
    return render_template('index.html')

if __name__ == '__main__':
    app.run(debug=True)
